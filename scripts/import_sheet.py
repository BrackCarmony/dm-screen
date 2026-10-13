"""Convert an xlsx export of "Brack's Battle Sheets" into seed JSON for the app.

Usage: python3 scripts/import_sheet.py path/to/export.xlsx [out_dir]
Needs openpyxl (pip install openpyxl).
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

SAVES = ["str", "dex", "con", "int", "wis", "cha"]
BATTLE_TABS = ["general", "army", "death slaad", "salamander", "trolls", "golem", "Copy of golem"]


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    try:
        f = float(str(v).strip())
        return int(f) if f.is_integer() else f
    except ValueError:
        return None


def text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-") or "x"


def bestiary(wb):
    ws = wb["Beastiary"]
    out, seen = [], set()
    for r in ws.iter_rows(min_row=4, min_col=2, max_col=19, values_only=True):
        name = text(r[0])
        if not name or name in seen:
            continue
        seen.add(name)
        ac = r[2]
        out.append({
            "name": name,
            "hp": num(r[1]),
            "ac": num(ac) if not hasattr(ac, "year") else None,
            "speed": text(r[3]),
            "saves": {k: num(r[4 + i]) for i, k in enumerate(SAVES)},
            "init": num(r[10]),
            "attacks": [text(a) for a in r[11:14] if text(a)],
            "cr": num(r[14]),
            "xp": num(r[15]),
            "special": text(r[16]),
            "url": text(r[17]),
        })
    return out


def encounters(wb):
    out = []
    ws = wb["Dungeon Rooms"]
    for r in ws.iter_rows(min_row=1, max_col=12, values_only=True):
        mons = [text(v) for v in r[2:12] if text(v)]
        if mons:
            num_ = text(r[1])
            label = text(r[0])
            out.append({"id": f"room-{len(out)}", "group": "Dungeon Rooms",
                        "name": f"Room {num_}" if num_ else "Room", "note": label, "monsters": mons})
    ws = wb["Combat Tables"]
    table = 0
    for r in ws.iter_rows(min_row=1, max_col=12, values_only=True):
        mons = [text(v) for v in r[3:12] if text(v)]
        if not mons:
            continue
        roll = num(r[2])
        if roll == 1 or table == 0:
            table += 1
        letter = chr(64 + table)
        name = f"Table {letter} · {roll}" if roll is not None else f"Table {letter} · extra"
        out.append({"id": f"table-{len(out)}", "group": "Combat Tables", "name": name, "note": "", "monsters": mons})
    return out


def shop(wb):
    ws = wb["Magic Shops Tool"]
    def table(name_col, mult_cell, target_cell, label):
        items = []
        for r in ws.iter_rows(min_row=7, min_col=name_col, max_col=name_col + 1, values_only=True):
            if text(r[0]) and num(r[1]) is not None:
                items.append({"name": text(r[0]), "price": num(r[1])})
        return {"id": slug(label), "name": label, "multiplier": num(ws[mult_cell].value) or 1,
                "goldTarget": num(ws[target_cell].value) or 0, "rolls": 3, "items": items}
    tables = [table(2, "D6", "B5", "Major items"), table(8, "J6", "H5", "Minor items")]
    saved, cur = [], None
    for r in ws.iter_rows(min_row=1, min_col=15, max_col=16, values_only=True):
        name, price = text(r[0]), num(r[1])
        if not name:
            continue
        if price is None and (cur is None or cur["items"]):
            cur = {"id": slug(name), "name": name.replace(" Current shop", ""), "items": []}
            saved.append(cur)
        elif cur is not None:
            cur["items"].append({"name": name, "price": price})
    return {"tables": tables, "saved": saved, "note": text(ws["A1"].value)}


def xp(wb):
    ws = wb["XP Chart"]
    levels, cr = [], []
    for r in ws.iter_rows(min_row=3, min_col=2, max_col=7, values_only=True):
        if num(r[0]) is not None:
            levels.append({"level": num(r[0]), "xp": num(r[1]), "prof": num(r[2])})
        if num(r[4]) is not None:
            cr.append({"cr": num(r[4]), "xp": num(r[5])})
    return {"levels": levels, "cr": cr}


def party(wb):
    ws = wb["DM Data"]
    out = []
    for r in ws.iter_rows(min_row=2, min_col=7, max_col=9, values_only=True):
        pid, name, normal = text(r[0]), text(r[1]) or text(r[2]), text(r[2])
        if pid and name:
            out.append({"id": pid, "name": name, "inParty": bool(normal)})
    return out


def battles(wb, roster):
    names = {p["id"]: p["name"] for p in roster}
    out = []
    for tab in BATTLE_TABS:
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        combatants = []
        inits = {}
        for r in ws.iter_rows(min_row=18, max_row=37, min_col=14, max_col=15, values_only=True):
            if text(r[0]) and num(r[1]) is not None:
                inits[text(r[0])] = num(r[1])
        for i, col in enumerate(range(2, 12)):
            name = text(ws.cell(3, col).value)
            if not name:
                continue
            dmg = [num(ws.cell(row, col).value) for row in range(10, 80)]
            combatants.append({
                "id": f"m{i + 1}", "kind": "monster", "name": name,
                "mask": text(ws.cell(9, col).value),
                "damage": [d for d in dmg if d is not None],
                "init": inits.get(f"M{i + 1:02d}"),
            })
        for code, val in inits.items():
            if code.startswith("P") and code in names:
                combatants.append({"id": code.lower(), "kind": "player", "name": names[code],
                                   "mask": "", "damage": [], "init": val})
        if any(c["kind"] == "monster" for c in combatants):
            out.append({"id": slug(tab), "name": tab, "save": text(ws["A4"].value).lower() or "con",
                        "partySize": num(ws["O40"].value) or 4, "partyDamage": None,
                        "round": 1, "turn": None, "combatants": combatants})
    return out


def main():
    src = Path(sys.argv[1])
    out = Path(sys.argv[2] if len(sys.argv) > 2 else Path(__file__).resolve().parent.parent / "data" / "seed")
    out.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(src, data_only=True)
    roster = party(wb)
    files = {
        "bestiary.json": bestiary(wb),
        "encounters.json": encounters(wb),
        "shop.json": shop(wb),
        "xp.json": xp(wb),
        "party.json": roster,
        "battles.json": battles(wb, roster),
    }
    for name, data in files.items():
        (out / name).write_text(json.dumps(data, indent=1, ensure_ascii=False) + "\n")
        print(name, len(data))


if __name__ == "__main__":
    main()
